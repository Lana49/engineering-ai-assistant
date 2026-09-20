"""Black-box tests for document and table ingestion.

The fixtures are generated in a TemporaryDirectory, so running this suite never
adds sample documents to ``data/raw`` and never depends on the real knowledge
base.  The tests intentionally exercise the public parser result (text, chunks,
and tables) instead of private implementation details.
"""

from __future__ import annotations

import csv
import importlib.util
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.parser import DocumentParser, parse_file  # noqa: E402
from core.table_extractor import extract_tables  # noqa: E402


def _module_available(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


def _table_text(table: Any) -> str:
    """Return one normalized searchable representation of a table object."""
    if isinstance(table, dict):
        headers = table.get("headers", [])
        rows = table.get("rows", [])
        raw = table.get("raw_text") or table.get("content") or ""
        title = table.get("title", "")
    else:
        headers = getattr(table, "headers", [])
        rows = getattr(table, "rows", [])
        raw = getattr(table, "raw_text", "")
        title = getattr(table, "title", "")

    lines = [str(title), str(raw)]
    lines.extend(" | ".join(str(cell) for cell in row) for row in [headers] if row)
    lines.extend(" | ".join(str(cell) for cell in row) for row in rows)
    return "\n".join(lines)


def _tables_from_parsed(parsed: dict[str, Any]) -> list[str]:
    """Collect unique tables regardless of whether they live on doc or chunks."""
    candidates: list[Any] = []
    candidates.extend(parsed.get("tables", []) or [])

    for chunk in parsed.get("chunks", []) or []:
        candidates.extend(chunk.get("tables", []) or [])
        metadata = chunk.get("metadata", {}) or {}
        candidates.extend(metadata.get("tables", []) or [])

    text = parsed.get("text", "") or ""
    if text:
        candidates.extend(extract_tables(text, doc_name=parsed.get("doc_name", "fixture"), min_rows=1))

    unique: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        normalized = " ".join(_table_text(candidate).split())
        if normalized and normalized not in seen:
            seen.add(normalized)
            unique.append(normalized)
    return unique


def _combined_chunk_text(parsed: dict[str, Any]) -> str:
    return "\n".join(str(chunk.get("text", "")) for chunk in parsed.get("chunks", []) or [])


class DocumentIngestionTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tempdir = tempfile.TemporaryDirectory(prefix="engineering-ingestion-")
        self.temp_path = Path(self._tempdir.name)

    def tearDown(self) -> None:
        self._tempdir.cleanup()

    def _assert_parsed(self, parsed: dict[str, Any], *tokens: str) -> None:
        self.assertTrue(parsed.get("metadata", {}).get("parsed"), parsed.get("metadata"))
        text = parsed.get("text", "")
        self.assertTrue(text.strip(), "Parser returned empty text")
        self.assertTrue(parsed.get("chunks"), "Parser returned no searchable chunks")
        for token in tokens:
            self.assertIn(token, text)

    def _assert_table_contains(self, parsed: dict[str, Any], *tokens: str) -> None:
        tables = _tables_from_parsed(parsed)
        self.assertTrue(tables, "No structured table was retained by the ingestion pipeline")
        self.assertTrue(
            any(all(token in table for token in tokens) for table in tables),
            f"No single table contains {tokens!r}; extracted tables: {tables!r}",
        )

    @unittest.skipUnless(
        _module_available("reportlab") and _module_available("pdfplumber"),
        "PDF table test requires reportlab and pdfplumber",
    )
    def test_pdf_preserves_text_and_ruled_table(self) -> None:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        path = self.temp_path / "climate-table.pdf"
        rows = [
            ["City", "Heating temperature", "Heating days"],
            ["Tomsk", "-8.4", "225"],
            ["Moscow", "-3.1", "214"],
        ]
        table = Table(rows, colWidths=[150, 140, 100])
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ]
            )
        )
        story = [
            Paragraph("Climate reference values", getSampleStyleSheet()["Heading2"]),
            Spacer(1, 8),
            table,
        ]
        SimpleDocTemplate(str(path), pagesize=A4).build(story)

        parsed = parse_file(path)

        self._assert_parsed(parsed, "Climate reference values", "Tomsk", "225")
        self._assert_table_contains(parsed, "City", "Heating days", "Tomsk", "225")

    @unittest.skipUnless(
        _module_available("pymupdf")
        and _module_available("PIL")
        and _module_available("pytesseract")
        and shutil.which("tesseract") is not None,
        "Mixed PDF test requires PyMuPDF, Pillow, pytesseract and Tesseract",
    )
    def test_mixed_pdf_uses_ocr_for_scanned_page(self) -> None:
        import pymupdf
        import pytesseract
        from PIL import Image, ImageDraw, ImageFont

        languages = set(pytesseract.get_languages(config=""))
        if not {"eng", "rus"}.issubset(languages):
            self.skipTest("The parser requests rus+eng, but those Tesseract languages are unavailable")

        image_path = self.temp_path / "scan.png"
        image = Image.new("RGB", (1800, 600), "white")
        draw = ImageDraw.Draw(image)

        font = None
        for candidate in (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/System/Library/Fonts/Helvetica.ttc",
        ):
            try:
                font = ImageFont.truetype(candidate, 100)
                break
            except OSError:
                continue
        if font is None:
            try:
                font = ImageFont.load_default(size=100)
            except TypeError:
                font = ImageFont.load_default()

        draw.text((80, 190), "SCANMARK 98765", fill="black", font=font)
        image.save(image_path)

        path = self.temp_path / "mixed.pdf"
        document = pymupdf.open()
        text_page = document.new_page()
        text_page.insert_text((72, 100), "TEXTMARK 12345", fontsize=24)
        scan_page = document.new_page(width=900, height=300)
        scan_page.insert_image(scan_page.rect, filename=str(image_path))
        document.save(path)
        document.close()

        parsed = parse_file(path)

        self._assert_parsed(parsed, "TEXTMARK", "12345", "SCANMARK", "98765")

    @unittest.skipUnless(_module_available("striprtf"), "RTF test requires striprtf")
    def test_rtf_cp1251_text_and_tabular_rows(self) -> None:
        path = self.temp_path / "climate-cp1251.rtf"
        rtf = (
            r"{\rtf1\ansi\ansicpg1251 "
            "Город\\tab Температура\\tab Дни\\par "
            "Томск\\tab -8,4\\tab 225\\par "
            "Москва\\tab -3,1\\tab 214\\par}"
        )
        path.write_bytes(rtf.encode("cp1251"))

        parsed = parse_file(path)

        self._assert_parsed(parsed, "Томск", "-8,4", "225")
        self.assertNotIn("�", parsed["text"], "cp1251 was decoded through UTF-8 replacement characters")
        self._assert_table_contains(parsed, "Город", "Температура", "Томск", "225")

    @unittest.skipUnless(_module_available("docx"), "DOCX test requires python-docx")
    def test_docx_keeps_body_order_and_separate_two_and_three_column_tables(self) -> None:
        from docx import Document

        path = self.temp_path / "ordered-tables.docx"
        document = Document()
        document.add_paragraph("BEFORE-TABLES")

        first = document.add_table(rows=3, cols=2)
        first_rows = [
            ["Parameter", "Value"],
            ["Air flow", "120"],
            ["Pressure", "15"],
        ]
        for row_index, row in enumerate(first_rows):
            for column_index, value in enumerate(row):
                first.cell(row_index, column_index).text = value

        document.add_paragraph("BETWEEN-TABLES")

        second = document.add_table(rows=3, cols=3)
        second_rows = [
            ["City", "Temperature", "Days"],
            ["Tomsk", "-8.4", "225"],
            ["Moscow", "-3.1", "214"],
        ]
        for row_index, row in enumerate(second_rows):
            for column_index, value in enumerate(row):
                second.cell(row_index, column_index).text = value

        document.add_paragraph("AFTER-TABLES")
        document.save(path)

        parsed = parse_file(path)
        self._assert_parsed(parsed, "BEFORE-TABLES", "BETWEEN-TABLES", "AFTER-TABLES")

        text = parsed["text"]
        ordered_tokens = ["BEFORE-TABLES", "Parameter", "BETWEEN-TABLES", "City", "AFTER-TABLES"]
        positions = [text.index(token) for token in ordered_tokens]
        self.assertEqual(positions, sorted(positions), f"DOCX body order was lost: {text!r}")

        tables = _tables_from_parsed(parsed)
        self.assertGreaterEqual(len(tables), 2, f"Expected two separate tables, got: {tables!r}")
        self.assertTrue(any(all(t in table for t in ("Parameter", "Value", "Air flow", "120")) for table in tables))
        self.assertTrue(any(all(t in table for t in ("City", "Temperature", "Tomsk", "225")) for table in tables))

    def test_csv_semicolon_utf8_bom_and_decimal_comma(self) -> None:
        path = self.temp_path / "climate.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerows(
                [
                    ["Город", "Температура", "Дни"],
                    ["Томск", "-8,4", "225"],
                    ["Москва", "-3,1", "214"],
                ]
            )

        parsed = parse_file(path)

        self._assert_parsed(parsed, "Томск", "-8,4", "225")
        self._assert_table_contains(parsed, "Город", "Температура", "Томск", "225")

    @unittest.skipUnless(_module_available("openpyxl"), "XLSX test requires openpyxl")
    def test_xlsx_reads_all_nonempty_sheets_as_separate_tables(self) -> None:
        from openpyxl import Workbook

        path = self.temp_path / "engineering.xlsx"
        workbook = Workbook()
        climate = workbook.active
        climate.title = "Climate"
        climate.append(["City", "Temperature", "Days"])
        climate.append(["Tomsk", -8.4, 225])

        materials = workbook.create_sheet("Materials")
        materials.append(["Material", "Conductivity"])
        materials.append(["Mineral wool", 0.045])
        workbook.save(path)

        parsed = parse_file(path)

        self._assert_parsed(parsed, "Climate", "Tomsk", "Materials", "Mineral wool")
        tables = _tables_from_parsed(parsed)
        self.assertGreaterEqual(len(tables), 2, f"Expected one table per nonempty sheet, got: {tables!r}")
        self.assertTrue(any(all(t in table for t in ("City", "Temperature", "Tomsk", "225")) for table in tables))
        self.assertTrue(any(all(t in table for t in ("Material", "Conductivity", "Mineral wool")) for table in tables))

    def test_long_table_chunks_repeat_header_and_never_split_a_row(self) -> None:
        path = self.temp_path / "long-table.csv"
        with path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["Item", "Value", "Unit"])
            for index in range(1, 81):
                writer.writerow([f"ITEM-{index:03d}", f"VALUE-{index:03d}", "kPa"])

        parsed = DocumentParser(chunk_size=300, chunk_overlap=50, min_chunk_size=50).parse_file(path)
        self._assert_parsed(parsed, "ITEM-001", "ITEM-080")

        chunks = [chunk.get("text", "") for chunk in parsed.get("chunks", [])]
        table_chunks = [chunk for chunk in chunks if "ITEM-" in chunk]
        self.assertGreater(len(table_chunks), 1, "Fixture should exercise table chunking, not one oversized chunk")
        for chunk in table_chunks:
            self.assertIn("Item", chunk, f"Table header was not repeated in chunk: {chunk!r}")
            self.assertIn("Value", chunk, f"Table header was not repeated in chunk: {chunk!r}")

        for target in ("001", "040", "080"):
            matching_chunks = [chunk for chunk in chunks if f"ITEM-{target}" in chunk]
            self.assertTrue(matching_chunks, f"Row ITEM-{target} disappeared during chunking")
            self.assertTrue(
                any(f"VALUE-{target}" in chunk and "kPa" in chunk for chunk in matching_chunks),
                f"Row ITEM-{target} was split across chunks: {matching_chunks!r}",
            )

        self._assert_table_contains(parsed, "Item", "Value", "ITEM-080", "VALUE-080")


if __name__ == "__main__":
    unittest.main(verbosity=2)
